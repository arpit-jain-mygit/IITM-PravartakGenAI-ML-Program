"""Helpers to write rich content (text, bullets, code, tables, images) into a Word document.

new_document() creates a plain A4 document (no template, logo or watermark); section() adds a heading;
Body(doc) writes text / bullets / code / tables / images into the document body."""
import copy
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_PT = 9.5


def new_document(title: str, subtitle: str = ""):
    """Plain A4 document with 2 cm margins and a title block - nothing else."""
    from docx import Document
    from docx.enum.section import WD_ORIENT  # noqa: F401  (kept for readers who want landscape pages)
    doc = Document()
    cp = doc.core_properties                  # file metadata: author and title of this report
    cp.author, cp.last_modified_by, cp.title, cp.comments = "Arpit Jain", "Arpit Jain", title, ""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(BODY_PT)
    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)
    t = doc.add_paragraph(); r = t.add_run(title); r.bold = True; r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
    if subtitle:
        s = doc.add_paragraph(); r = s.add_run(subtitle); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return doc


def section(doc, title: str, level: int = 1):
    """Section heading in the document body (level 1 = task, level 2 = sub-task)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.keep_with_next = Pt(10 if level == 1 else 6), Pt(3), True
    r = p.add_run(title); r.bold = True; r.font.size = Pt(13 if level == 1 else 11)
    r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)
    return p


class Cell:
    fit_cm = 14.6                       # usable width inside a template answer box

    def __init__(self, cell, width_cm=15.6):
        self.c = cell
        self.width = width_cm
        # remove the "[Insert your response here...]" placeholder paragraphs
        for p in list(cell.paragraphs):
            p._p.getparent().remove(p._p)

    def _para(self, space_after=3):
        p = self.c.add_paragraph()
        pf = p.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing = Pt(space_after), Pt(0), 1.1
        return p

    def _runs(self, p, text, size=BODY_PT, bold=False):
        # very small markup: **bold** segments
        parts = text.split("**")
        for i, part in enumerate(parts):
            if part:
                r = p.add_run(part)
                r.font.size = Pt(size)
                r.bold = bold or (i % 2 == 1)
        return p

    def heading(self, text):
        p = self._para(space_after=2)
        p.paragraph_format.space_before = Pt(5)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)

    def text(self, text, size=BODY_PT):
        self._runs(self._para(), text, size)

    def bullets(self, items, size=BODY_PT):
        for t in items:
            p = self._para(space_after=1.5)
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.3)
            self._runs(p, "• " + t, size)

    def code(self, src, caption=None):
        if caption:
            p = self._para(space_after=1); r = p.add_run(caption); r.italic = True; r.font.size = Pt(8)
        lines = src.strip("\n").split("\n")
        p = self._para(space_after=5)
        p.paragraph_format.line_spacing = 1.0
        # light grey shading
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F4F7")
        p._p.get_or_add_pPr().append(shd)
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = "Courier New"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            r.font.size = Pt(7.5)
            if i < len(lines) - 1:
                r.add_break()

    def image(self, path, width_cm=None, caption=None):
        p = self._para(space_after=1); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm or self.width))
        if caption:
            q = self._para(space_after=5); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(caption); r.italic = True; r.font.size = Pt(8)

    def table(self, rows, col_widths_cm=None, size=8, header_fill="1F3B57", bold_row=None):
        t = self.c.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        if col_widths_cm:
            # scale columns so the table always fits inside the template's answer box
            total = sum(col_widths_cm)
            col_widths_cm = [w * self.fit_cm / total for w in col_widths_cm]
            t.autofit = False
            # explicit grid widths so Word and LibreOffice both respect the column sizes
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
                gc.set(qn("w:w"), str(int(col_widths_cm[j] * 567)))
            for j, col in enumerate(t.columns):
                col.width = Cm(col_widths_cm[j])
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(str(val)); r.font.size = Pt(size)
                if i == 0:
                    r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), header_fill)
                    cell._tc.get_or_add_tcPr().append(shd)
                elif bold_row is not None and i == bold_row:
                    r.bold = True
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "DFF0E4")
                    cell._tc.get_or_add_tcPr().append(shd)
                if col_widths_cm:
                    cell.width = Cm(col_widths_cm[j])
        if isinstance(self, Body):
            self._para(space_after=4)        # body tables need an explicit spacer paragraph
        else:
            # python-docx already adds an empty paragraph after a table inside a cell; make it a small spacer
            spacer = self.c.paragraphs[-1].paragraph_format
            spacer.space_after, spacer.space_before, spacer.line_spacing = Pt(2), Pt(0), 0.8
        return t


class Body(Cell):
    """Writes into the document body instead of a table cell (no placeholder clearing)."""
    fit_cm = 16.8

    def __init__(self, doc, width_cm=16.0):
        self.c = doc
        self.width = width_cm
