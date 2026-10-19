"""Helpers to write rich content (text, bullets, code, tables, images) into a template table cell."""
import copy
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_PT = 9.5


class Cell:
    def __init__(self, cell, width_cm=15.6):
        self.c = cell
        self.width = width_cm
        # remove the "[Insert your response here...]" placeholder paragraphs
        for p in list(cell.paragraphs):
            p._p.getparent().remove(p._p)

    def _para(self, space_after=3):
        p = self.c.add_paragraph()
        pf = p.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing = Pt(space_after), Pt(0), 1.1
        return p

    def _runs(self, p, text, size=BODY_PT, bold=False):
        # very small markup: **bold** segments
        parts = text.split("**")
        for i, part in enumerate(parts):
            if part:
                r = p.add_run(part)
                r.font.size = Pt(size)
                r.bold = bold or (i % 2 == 1)
        return p

    def heading(self, text):
        p = self._para(space_after=2)
        p.paragraph_format.space_before = Pt(5)
        r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x57)

    def text(self, text, size=BODY_PT):
        self._runs(self._para(), text, size)

    def bullets(self, items, size=BODY_PT):
        for t in items:
            p = self._para(space_after=1.5)
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.3)
            self._runs(p, "• " + t, size)

    def code(self, src, caption=None):
        if caption:
            p = self._para(space_after=1); r = p.add_run(caption); r.italic = True; r.font.size = Pt(8)
        lines = src.strip("\n").split("\n")
        p = self._para(space_after=5)
        p.paragraph_format.line_spacing = 1.0
        # light grey shading
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F4F7")
        p._p.get_or_add_pPr().append(shd)
        for i, line in enumerate(lines):
            r = p.add_run(line)
            r.font.name = "Courier New"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            r.font.size = Pt(7.5)
            if i < len(lines) - 1:
                r.add_break()

    def image(self, path, width_cm=None, caption=None):
        p = self._para(space_after=1); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(width_cm or self.width))
        if caption:
            q = self._para(space_after=5); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = q.add_run(caption); r.italic = True; r.font.size = Pt(8)

    def table(self, rows, col_widths_cm=None, size=8, header_fill="1F3B57", bold_row=None):
        t = self.c.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        if col_widths_cm:
            # scale columns so the table always fits inside the template's answer box
            total = sum(col_widths_cm)
            col_widths_cm = [w * 14.6 / total for w in col_widths_cm]
            t.autofit = False
            # explicit grid widths so Word and LibreOffice both respect the column sizes
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
                gc.set(qn("w:w"), str(int(col_widths_cm[j] * 567)))
            for j, col in enumerate(t.columns):
                col.width = Cm(col_widths_cm[j])
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(str(val)); r.font.size = Pt(size)
                if i == 0:
                    r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), header_fill)
                    cell._tc.get_or_add_tcPr().append(shd)
                elif bold_row is not None and i == bold_row:
                    r.bold = True
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "DFF0E4")
                    cell._tc.get_or_add_tcPr().append(shd)
                if col_widths_cm:
                    cell.width = Cm(col_widths_cm[j])
        # python-docx already adds an empty paragraph after a table inside a cell; make it a small spacer
        spacer = self.c.paragraphs[-1].paragraph_format
        spacer.space_after, spacer.space_before, spacer.line_spacing = Pt(2), Pt(0), 0.8
        return t
